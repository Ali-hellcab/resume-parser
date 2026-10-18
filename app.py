"""
Multimodal Resume Parsing Agent (Streamlit)
===========================================

Pipeline:  upload -> extract raw text -> Groq LLM -> Pydantic validation -> editable form

------------------------------------------------------------------------------
DEPLOYING TO STREAMLIT COMMUNITY CLOUD
------------------------------------------------------------------------------
1. Create a GitHub repo (public, or private with Streamlit access granted) and
   commit these files at the repo root:
       app.py
       requirements.txt
2. Get a free Groq API key at https://console.groq.com/keys
3. For local development, create `.streamlit/secrets.toml` with:
       GROQ_API_KEY = "gsk_your_key_here"
   This file is gitignored; do not add the key directly to this source file.
4. Go to https://share.streamlit.io -> "Create app" -> "Deploy a public app
   from GitHub". Choose your repo, branch, and set "Main file path" to app.py.
5. Before deploying, open "Advanced settings" and:
     - choose Python 3.11 (best wheel support for torch/easyocr)
     - paste into the Secrets box:
           GROQ_API_KEY = "gsk_your_key_here"
6. Click Deploy. The first build takes several minutes (torch + easyocr).
   Never commit your API key to GitHub; use Secrets only.

MEMORY NOTES (Community Cloud gives roughly 1GB RAM):
 - requirements.txt installs CPU-only torch (no CUDA).
 - Heavy models (EasyOCR, Whisper) are loaded lazily, only when a matching
   file type is uploaded, and are freed immediately after use (gc.collect()).
 - Whisper uses the "tiny" int8 model; images are downscaled before OCR;
   PDFs are read page by page with caches flushed; text sent to the LLM is capped.
 - Only one heavy model is ever resident in memory at a time.
------------------------------------------------------------------------------
"""

import gc
import json
import os
import tempfile
from io import BytesIO
from typing import List

import streamlit as st
from streamlit.errors import StreamlitSecretNotFoundError
from pydantic import BaseModel, ConfigDict, Field, ValidationError

# ----------------------------- Configuration --------------------------------
GROQ_MODEL = "openai/gpt-oss-120b"
WHISPER_MODEL = "tiny"                    # "tiny" ~150MB RAM; "base" ~300MB
MAX_FILE_MB = 15
MAX_TEXT_CHARS = 12000                    # cap on text sent to the LLM
MAX_IMAGE_SIDE = 1600                     # px, downscale before OCR
MAX_PDF_PAGES = 15

st.set_page_config(page_title="Resume Parsing Agent", page_icon="📄", layout="centered")


# ----------------------- Strict structured output schema ---------------------
class ResumeData(BaseModel):
    """Exact keys required by the spec; extra keys are rejected."""
    model_config = ConfigDict(populate_by_name=True, extra="forbid")

    full_name: str = Field("", alias="Full Name")
    email: str = Field("", alias="Email")
    phone_number: str = Field("", alias="Phone Number")
    education: List[str] = Field(default_factory=list, alias="Education")
    qualifications: List[str] = Field(default_factory=list, alias="Qualifications")


SYSTEM_PROMPT = f"""You are a resume parsing agent. Extract information from the raw text of a
resume (it may be noisy OCR or a spoken transcript). Respond with ONLY a JSON object
with EXACTLY these keys and no others:
  "Full Name": string
  "Email": string
  "Phone Number": string
  "Education": array of strings (one entry per degree/school, e.g. "B.Sc. Computer Science, MIT, 2020")
  "Qualifications": array of strings (skills, certifications, key experience highlights)
Use "" or [] for anything not present. Never invent information.
Schema reference: {json.dumps(ResumeData.model_json_schema(by_alias=True))}"""


# ------------------------------ Extractors -----------------------------------
def extract_pdf(data: bytes) -> str:
    import pdfplumber
    parts = []
    with pdfplumber.open(BytesIO(data)) as pdf:
        for page in pdf.pages[:MAX_PDF_PAGES]:
            parts.append(page.extract_text() or "")
            page.flush_cache()  # release per-page cache to keep memory flat
    return "\n".join(parts)


def extract_docx(data: bytes) -> str:
    from docx import Document
    doc = Document(BytesIO(data))
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:  # resumes often use tables for layout
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_image(data: bytes) -> str:
    import easyocr
    import numpy as np
    from PIL import Image

    img = Image.open(BytesIO(data)).convert("RGB")
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))  # cap pixels -> cap RAM
    arr = np.array(img)
    del img

    # Lazy load; CPU only; quantize to shrink the model footprint
    reader = easyocr.Reader(["en"], gpu=False, quantize=True, verbose=False)
    try:
        lines = reader.readtext(arr, detail=0, paragraph=True)
    finally:
        del reader, arr
        gc.collect()
    return "\n".join(lines)


def extract_audio(data: bytes, suffix: str) -> str:
    from faster_whisper import WhisperModel

    # faster-whisper decodes via bundled PyAV; write to a temp file for reliability
    with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
        tmp.write(data)
        path = tmp.name
    model = WhisperModel(WHISPER_MODEL, device="cpu", compute_type="int8", cpu_threads=2)
    try:
        segments, _ = model.transcribe(path, beam_size=1, vad_filter=True)
        text = " ".join(seg.text.strip() for seg in segments)  # generator: streamed
    finally:
        del model
        os.remove(path)
        gc.collect()
    return text


def extract_text(name: str, data: bytes) -> str:
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return extract_pdf(data)
    if ext == "docx":
        return extract_docx(data)
    if ext in ("png", "jpg", "jpeg"):
        return extract_image(data)
    if ext in ("mp3", "wav"):
        return extract_audio(data, "." + ext)
    raise ValueError(f"Unsupported file type: .{ext}")


# ------------------------------ LLM agent ------------------------------------
def get_api_key() -> str:
    try:
        return st.secrets["GROQ_API_KEY"]
    except (KeyError, StreamlitSecretNotFoundError):
        return os.environ.get("GROQ_API_KEY", "")


def parse_with_llm(raw_text: str, api_key: str) -> ResumeData:
    from groq import Groq

    client = Groq(api_key=api_key)
    messages = [
        {"role": "system", "content": SYSTEM_PROMPT},
        {"role": "user", "content": raw_text[:MAX_TEXT_CHARS]},
    ]
    last_err = None
    for _ in range(2):  # one retry, feeding the validation error back to the model
        resp = client.chat.completions.create(
            model=GROQ_MODEL,
            messages=messages,
            temperature=0,
            response_format={"type": "json_object"},  # Groq JSON mode
        )
        content = resp.choices[0].message.content
        try:
            return ResumeData.model_validate_json(content)  # strict validation
        except ValidationError as e:
            last_err = e
            messages += [
                {"role": "assistant", "content": content},
                {"role": "user", "content": f"Invalid output: {e}. Return corrected JSON with exact keys only."},
            ]
    raise RuntimeError(f"LLM output failed schema validation: {last_err}")


# --------------------------------- UI ----------------------------------------
st.title("📄 Multimodal Resume Parsing Agent")
st.caption("Upload a PDF, DOCX, image, or audio resume. Review and edit the extracted fields.")

api_key = get_api_key()

uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "png", "jpg", "jpeg", "mp3", "wav"])

if uploaded and st.button("Extract information", type="primary"):
    if not api_key:
        st.error("GROQ_API_KEY is not configured. Add it to Streamlit secrets or set it as an environment variable.")
    elif uploaded.size > MAX_FILE_MB * 1024 * 1024:
        st.error(f"File too large (limit {MAX_FILE_MB} MB).")
    else:
        try:
            with st.spinner("Extracting text..."):
                raw = extract_text(uploaded.name, uploaded.getvalue())
            if not raw.strip():
                st.warning("No text could be extracted from this file.")
            else:
                with st.spinner("Parsing with LLM..."):
                    result = parse_with_llm(raw, api_key)
                # Seed the editable widgets via session_state
                st.session_state.update(
                    f_name=result.full_name,
                    f_email=result.email,
                    f_phone=result.phone_number,
                    f_edu="\n".join(result.education),
                    f_qual="\n".join(result.qualifications),
                    raw_text=raw,
                )
        except Exception as e:  # surface any failure cleanly
            st.error(f"Failed: {e}")
        finally:
            gc.collect()

if "f_name" in st.session_state:
    st.subheader("Extracted information (editable)")
    st.text_input("Full Name", key="f_name")
    st.text_input("Email", key="f_email")
    st.text_input("Phone Number", key="f_phone")
    st.text_area("Education (one per line)", key="f_edu", height=120)
    st.text_area("Qualifications (one per line)", key="f_qual", height=160)

    final = {
        "Full Name": st.session_state.f_name,
        "Email": st.session_state.f_email,
        "Phone Number": st.session_state.f_phone,
        "Education": [l for l in st.session_state.f_edu.splitlines() if l.strip()],
        "Qualifications": [l for l in st.session_state.f_qual.splitlines() if l.strip()],
    }
    st.download_button("Download JSON", json.dumps(final, indent=2),
                       file_name="resume.json", mime="application/json")
    with st.expander("Raw extracted text"):
        st.text(st.session_state.get("raw_text", "")[:5000])
